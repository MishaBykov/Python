import openpyxl

id_columns = 'R9C2', 'R9C2'
compare_column = 'R9C9', 'R9C9'

wb = openpyxl.load_workbook(filename='оборотно-сальдовая_июнь.xlsx')
sheet = wb['Report']
vals = [
    sheet.cell(9, 2).value,
    '|',
    sheet.cell(9, 9).value
]

print(vals)

# import xlrd, xlwt

# открываем файл
# rb = xlrd.open_workbook('оборотно-сальдовая_июнь.xlsx', formatting_info=True)

# выбираем активный лист
# sheet = rb.sheet_by_index(0)
